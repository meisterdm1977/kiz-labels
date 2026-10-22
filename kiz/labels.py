"""Reading label layouts exported from 1С: codes, product fields and print-ready pages."""
import io
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageFont, ImageOps
from pylibdmtx.pylibdmtx import decode as dmtx_decode
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject

from . import config

GS = "\x1d"
MM = 72 / 25.4
CODE_RE = re.compile(r"^01(\d{14})21(.+?)(?:\x1d|$)")
TEXT_CODE_RE = re.compile(r"\(01\)\s*(\d{14})\s*\(21\)\s*(\S+)")
EAN_RE = re.compile(r"^\d{13}$")
FIELD_LABELS = {"артикул": "article", "размер": "size", "цвет": "color"}


class LayoutError(Exception):
    pass


@dataclass
class Label:
    index: int
    ean: str = ""
    article: str = ""
    size: str = ""
    color: str = ""
    code: str = ""
    errors: list = field(default_factory=list)

    @property
    def code_prefix(self):
        return code_prefix(self.code)


_RU = "йцукенгшщзхъфывапролджэячсмитьбюё" + "ЙЦУКЕНГШЩЗХЪФЫВАПРОЛДЖЭЯЧСМИТЬБЮЁ"
_EN = "qwertyuiop[]asdfghjkl;'zxcvbnm,.`" + "QWERTYUIOP{}ASDFGHJKL:\"ZXCVBNM<>~"
_FROM_RU_LAYOUT = str.maketrans(_RU, _EN)


def from_ru_layout(text):
    """A keyboard-emulating scanner on a Russian layout types Cyrillic letters; map them back to the keys."""
    return text.translate(_FROM_RU_LAYOUT)


def code_prefix(code):
    """GTIN + serial: what identifies a code even when a scanner drops the GS separators."""
    m = CODE_RE.match(code)
    if not m:
        return code.replace(GS, "")[:31]
    return f"01{m.group(1)}21{m.group(2)}"


def printable_code(code):
    return code.replace(GS, " ")


def decode_datamatrix(data):
    img = Image.open(io.BytesIO(data)).convert("L")
    w, h = img.size
    if min(w, h) < 40:
        return None
    # Images in the layout have no quiet zone and are scaled up a lot; normalise both.
    target = 240
    if max(w, h) > target:
        img = img.resize((target, round(h * target / w)), Image.NEAREST)
    img = ImageOps.expand(img, border=max(img.size) // 8, fill=255)
    found = dmtx_decode(img, max_count=1, timeout=3000)
    if not found:
        return None
    return found[0].data.decode("latin-1")


def _blocks(ws):
    breaks = sorted(b.id for b in ws.row_breaks.brk)
    if not breaks:
        raise LayoutError("В файле нет разбивки на страницы: не понятно, где кончается одна этикетка.")
    start, blocks = 1, []
    for end in breaks:
        if end >= start:
            blocks.append((start, end))
            start = end + 1
    if start <= ws.max_row:
        blocks.append((start, ws.max_row))
    return blocks


def _fill_fields(label, cells):
    rows = {}
    for (row, col), value in cells.items():
        rows.setdefault(row, []).append((col, value))
    for row in rows.values():
        row.sort()
        for i, (_, value) in enumerate(row):
            key = FIELD_LABELS.get(value.strip().rstrip(":").strip().lower())
            if key and i + 1 < len(row):
                setattr(label, key, row[i + 1][1].strip())


def read_workbook(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    blocks = _blocks(ws)
    images = {}
    for im in ws._images:
        images.setdefault(im.anchor._from.row + 1, []).append(im)

    labels = []
    for index, (start, end) in enumerate(blocks):
        cells = {}
        for row in ws.iter_rows(min_row=start, max_row=end):
            for c in row:
                if isinstance(c.value, str) and c.value.strip():
                    cells[(c.row, c.column)] = c.value
                elif isinstance(c.value, (int, float)) and len(str(int(c.value))) == 13:
                    cells[(c.row, c.column)] = str(int(c.value))
        block_images = [im for r in range(start, end + 1) for im in images.get(r, [])]
        if not cells and not block_images:
            continue

        label = Label(index=len(labels))
        _fill_fields(label, cells)
        label.ean = next((v.strip() for v in cells.values() if EAN_RE.match(v.strip())), "")
        text_code = next((TEXT_CODE_RE.search(v) for v in cells.values() if TEXT_CODE_RE.search(v)), None)

        block_images.sort(key=lambda im: -(im.width * im.height))
        for im in block_images:
            decoded = decode_datamatrix(im._data())
            if decoded and decoded.startswith("01"):
                label.code = decoded
                break

        if not label.code:
            label.errors.append("не удалось прочитать DataMatrix")
        else:
            m = CODE_RE.match(label.code)
            if not m:
                label.errors.append("DataMatrix не похож на код маркировки")
            elif GS not in label.code:
                label.errors.append("в коде нет крипто-хвоста")
            elif text_code and (text_code.group(1), text_code.group(2)) != (m.group(1), m.group(2)):
                label.errors.append("код в картинке не совпадает с текстом на этикетке")
        if not label.ean:
            label.errors.append("не найден штрихкод EAN-13")
        if not label.article:
            label.errors.append("не найден артикул")
        labels.append(label)
    if not labels:
        raise LayoutError("В файле не найдено ни одной этикетки.")
    return labels


_FONT_FILES = {
    False: ["/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"],
    True: ["/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
           "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"],
}
_fonts = {}


def _measure_font(bold, size):
    key = (bold, round(size * 2))
    if key not in _fonts:
        path = next((f for f in _FONT_FILES[bold] if Path(f).exists()), None)
        _fonts[key] = ImageFont.truetype(path, size) if path else ImageFont.load_default()
    return _fonts[key]


def _wrap(text, font, width):
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if line and font.getlength(candidate) > width:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


FIELD_NAMES = ("пол", "артикул", "размер", "цвет", "состав")
LAYOUT_VERSION = "2"  # bump when the printed page changes, so old uploads are rendered again


def _merged_range(ws, cell):
    for merged in ws.merged_cells.ranges:
        if cell.coordinate in merged:
            return merged
    return None


def _digit_width(wb):
    # Column width is counted in widest digits of the workbook's default font, plus 5 px padding.
    default = wb._fonts[0]
    return _measure_font(bool(default.b), (default.sz or 11) * 96 / 72).getlength("0")


def _field_rows(ws):
    """(name cell, its range, value cell, its range) for Пол/Артикул/Размер/Цвет/Состав lines."""
    for row in ws.iter_rows():
        for label in row:
            text = label.value.strip().lower() if isinstance(label.value, str) else ""
            if not text.endswith(":") or text.rstrip(":").strip() not in FIELD_NAMES:
                continue
            label_range = _merged_range(ws, label)
            label_end = label_range.max_col if label_range else label.column
            value = next((c for c in row if c.column > label_end and isinstance(c.value, str) and c.value.strip()), None)
            if value is not None:
                yield label, label_range, value, _merged_range(ws, value)


def _widen_fields(wb):
    """In the 1С layout a field name («Артикул:») takes half the width and its value the other half,
    so a long article or composition is cut. When a value does not fit, the lines become
    «Артикул: value» across the whole width, starting where the name starts."""
    digit = _digit_width(wb)
    for ws in wb.worksheets:
        fields = list(_field_rows(ws))
        if not any(merged and not _cell_fits(ws, merged, value, digit, float(value.font.sz or 11))
                   for _, _, value, merged in fields):
            continue
        for label, label_range, value, value_range in fields:
            top = min(label_range.min_row if label_range else label.row, value_range.min_row if value_range else value.row)
            bottom = max(label_range.max_row if label_range else label.row, value_range.max_row if value_range else value.row)
            first = label_range.min_col if label_range else label.column
            last = value_range.max_col if value_range else value.column
            others = [ws.cell(r, c).value for r in range(top, bottom + 1) for c in range(first, last + 1)
                      if (r, c) not in ((label.row, label.column), (value.row, value.column))]
            if any(v not in (None, "") for v in others):
                continue  # something else lives there; leave this line as 1С made it
            text = f"{label.value.strip()} {value.value.strip()}"
            style = value._style
            for merged in (label_range, value_range):
                if merged:
                    ws.unmerge_cells(str(merged))
            label.value = value.value = None
            cell = ws.cell(top, first, text)
            cell._style = style
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.merge_cells(start_row=top, start_column=first, end_row=bottom, end_column=last)


def _fit_texts(wb):
    """Text in a merged cell of the 1С layout is cut off when it is longer than the cell (a long
    article, a long composition). Shrink the font just enough to show it whole; text that already
    fits is left alone."""
    digit = _digit_width(wb)
    for ws in wb.worksheets:
        for merged in ws.merged_cells.ranges:
            cell = ws.cell(merged.min_row, merged.min_col)
            if isinstance(cell.value, str) and cell.value.strip():
                _fit_cell(ws, merged, cell, digit)


def _cell_fits(ws, merged, cell, digit, size):
    # A little narrower than measured: LibreOffice's text is not exactly as wide as ours.
    width = (sum(round((ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 8.43) * digit) + 5
                 for c in range(merged.min_col, merged.max_col + 1)) * 0.75 - 3) * 0.95
    height = sum(ws.row_dimensions[r].height or 15 for r in range(merged.min_row, merged.max_row + 1))
    text = cell.value.strip()
    font = _measure_font(bool(cell.font.b), size)
    if not cell.alignment.wrap_text:
        return all(font.getlength(line) <= width for line in text.split("\n"))
    lines = _wrap(text, font, width)
    if len(lines) <= len(text.split("\n")):
        return True
    return len(lines) * size * 1.35 <= height


def _fit_cell(ws, merged, cell, digit):
    size = original = float(cell.font.sz or 11)
    while size > 4.5 and not _cell_fits(ws, merged, cell, digit, size):
        size -= 0.5
    if size != original:
        cell.font = Font(**{**_font_kwargs(cell.font), "sz": size})


def _font_kwargs(font):
    return {k: getattr(font, k) for k in ("name", "b", "i", "u", "strike", "color", "vertAlign", "charset",
                                          "outline", "shadow", "condense", "extend", "family", "scheme")}


def render_pdf(xlsx_path, out_pdf, expected_pages):
    """Convert the layout to PDF and cut each page down to the label size."""
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "labels.xlsx"
        wb = openpyxl.load_workbook(xlsx_path)
        _widen_fields(wb)
        _fit_texts(wb)
        wb.save(src)
        result = subprocess.run(
            [config.SOFFICE, f"-env:UserInstallation=file://{tmp}/profile",
             "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
            capture_output=True, timeout=900,
        )
        pdf = Path(tmp) / "labels.pdf"
        if result.returncode != 0 or not pdf.exists():
            raise LayoutError("Не удалось преобразовать файл в этикетки для печати.")

        reader = PdfReader(pdf)
        pages = list(reader.pages)
        # A layout can end with an empty page after the last break.
        while len(pages) > expected_pages and not (pages[-1].extract_text() or "").strip():
            pages.pop()
        if len(pages) != expected_pages:
            raise LayoutError(
                f"Этикеток в файле {expected_pages}, а страниц для печати получилось {len(pages)}."
            )

        width, height = config.LABEL_WIDTH_MM * MM, config.LABEL_HEIGHT_MM * MM
        writer = PdfWriter()
        for page in pages:
            box = page.mediabox
            left, top = float(box.left), float(box.top)
            page.mediabox.lower_left = (left, top - height)
            page.mediabox.upper_right = (left + width, top)
            page.cropbox.lower_left = page.mediabox.lower_left
            page.cropbox.upper_right = page.mediabox.upper_right
            _keep_own_images(page)
            writer.add_page(page)
        with open(out_pdf, "wb") as f:
            writer.write(f)


def _keep_own_images(page):
    """LibreOffice gives every page the images of the whole file; keep only the ones this page draws."""
    resources = page.get("/Resources")
    if not resources or "/XObject" not in resources:
        return
    used = {m.decode() for m in re.findall(rb"/([^\s/\[\]()<>{}%]+)\s+Do\b", page.get_contents().get_data())}
    xobjects = resources["/XObject"].get_object()
    own = DictionaryObject({NameObject(k): v for k, v in xobjects.items() if k.lstrip("/") in used})
    new_resources = DictionaryObject(resources.get_object())
    new_resources[NameObject("/XObject")] = own
    page[NameObject("/Resources")] = new_resources


def build_print_pdf(items, out):
    """items: (labels_pdf_path, page_no) in print order."""
    readers, writer = {}, PdfWriter()
    for path, page_no in items:
        if path not in readers:
            readers[path] = PdfReader(path)
        writer.add_page(readers[path].pages[page_no])
    writer.write(out)
