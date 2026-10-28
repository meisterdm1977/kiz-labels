import asyncio
import hashlib
import io
import logging
import re
import secrets
import subprocess
import tempfile
import time
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from psycopg.errors import UniqueViolation
from starlette.exceptions import HTTPException as StarletteHTTPException
from starlette.middleware.sessions import SessionMiddleware

from . import auth, config, db, nomenclature, service, stickers
from .auth import AdminRequired, DeviceRequired, require_admin, require_device
from .badges import badge_pdf
from .labels import build_print_pdf, printable_code
from .small_label import SmallLabelError
from .small_label import build_pdf as build_small_pdf

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("kiz")

HERE = Path(__file__).parent

async def _cleanup_loop():
    while True:
        try:
            await asyncio.to_thread(service.cleanup)
        except Exception:
            log.exception("cleanup failed")
        await asyncio.sleep(6 * 3600)


@asynccontextmanager
async def lifespan(app):
    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    db.init()
    auth.ensure_admin()
    service.recover_interrupted_uploads()
    service.rerender_old_layouts()
    task = asyncio.create_task(_cleanup_loop())
    yield
    task.cancel()
    db.pool.close()


app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
app.add_middleware(
    SessionMiddleware,
    secret_key=config.SECRET_KEY,
    session_cookie="kiz_session",
    max_age=60 * 60 * 24 * 30,
    same_site="lax",
)
app.mount("/static", StaticFiles(directory=HERE / "static"), name="static")
templates = Jinja2Templates(directory=HERE / "templates")


def _local(value):
    if not value:
        return ""
    return value.astimezone(config.TIMEZONE).strftime("%d.%m.%Y %H:%M")


templates.env.filters["local"] = _local
templates.env.filters["code"] = printable_code
templates.env.globals["role_names"] = auth.ROLE_NAMES
# Label sizes for the size switch: the 1С page as is, or the small label drawn by the service.
SIZES = {"50x104": (50, 104), "58x40": (58, 40)}
DEFAULT_SIZE = "50x104"
templates.env.globals["sizes"] = SIZES
# Changes with app.css/app.js, so browsers fetch the new files after an update instead of a cached copy.
templates.env.globals["static_version"] = hashlib.md5(
    b"".join((HERE / "static" / f).read_bytes() for f in ("app.css", "app.js"))).hexdigest()[:10]


def _size(value):
    return value if value in SIZES else DEFAULT_SIZE


@app.exception_handler(StarletteHTTPException)
async def _http_error(request, exc):
    texts = {403: "Недостаточно прав для этой страницы.", 404: "Такой страницы нет."}
    return templates.TemplateResponse(
        request, "error.html", {"flash": None, "text": texts.get(exc.status_code, str(exc.detail))},
        status_code=exc.status_code,
    )


@app.exception_handler(RequestValidationError)
async def _bad_request(request, exc):
    return templates.TemplateResponse(
        request, "error.html", {"flash": None, "text": "Неверный адрес или данные формы."},
        status_code=400,
    )


@app.exception_handler(DeviceRequired)
async def _device_required(request, exc):
    return RedirectResponse("/device", status_code=303)


@app.exception_handler(AdminRequired)
async def _admin_required(request, exc):
    return RedirectResponse(f"/admin?next={request.url.path}", status_code=303)


def render(request, name, **context):
    flash = request.session.pop("flash", None)
    return templates.TemplateResponse(request, name, {
        "device": getattr(request.state, "device", None),
        "admin": getattr(request.state, "admin", None),
        "flash": flash,
        **context,
    })


def redirect(url, request=None, message=None, error=False):
    if request is not None and message:
        request.session["flash"] = {"text": message, "error": error}
    return RedirectResponse(url, status_code=303)


def _date(value):
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


def _int(value):
    try:
        return int(value) if value not in (None, "") else None
    except ValueError:
        return None


# ---------- computer approval and administrator ----------

def _safe_next(url, default):
    return url if url.startswith("/") and not url.startswith("//") else default


def _check_admin_password(request, password):
    """The administrator found by password, or an error text. Wrong passwords count towards a block."""
    key = auth.client_key(request)
    if auth.login_throttle.blocked(key):
        return None, "Слишком много неудачных попыток. Попробуйте через 15 минут."
    with db.pool.connection() as conn:
        admin = auth.admin_by_password(conn, password)
    if not admin:
        auth.login_throttle.fail(key)
        return None, "Неверный пароль администратора"
    auth.login_throttle.reset(key)
    return admin, None


@app.get("/device", response_class=HTMLResponse)
def device_page(request: Request):
    if auth.current_device(request):
        return redirect("/pack")
    return render(request, "device.html")


@app.post("/device")
def device_approve(request: Request, name: str = Form(""), password: str = Form(...)):
    admin, error = _check_admin_password(request, password)
    if not admin:
        return redirect("/device", request, error, error=True)
    with db.tx() as conn:
        token = auth.approve_device(conn, name.strip()[:60] or "Компьютер", admin, auth.client_key(request))
    response = redirect("/pack", request, "Компьютер допущен к работе")
    response.set_cookie(auth.DEVICE_COOKIE, token, max_age=auth.DEVICE_COOKIE_DAYS * 86400, httponly=True,
                        samesite="lax", secure=request.url.scheme == "https")
    return response


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request, next: str = "/stock"):
    require_device(request)
    if request.state.admin:
        return redirect(_safe_next(next, "/stock"))
    return render(request, "admin_login.html", next=next)


@app.post("/admin")
def admin_login(request: Request, password: str = Form(...), next: str = Form("/stock")):
    require_device(request)
    admin, error = _check_admin_password(request, password)
    if not admin:
        return redirect(f"/admin?next={next}", request, error, error=True)
    request.session["admin_id"] = admin["id"]
    request.session["admin_seen"] = time.time()
    return redirect(_safe_next(next, "/stock"))


@app.post("/admin/logout")
def admin_logout(request: Request):
    request.session.pop("admin_id", None)
    request.session.pop("admin_seen", None)
    return redirect("/pack")


@app.get("/login")
def old_login():
    return redirect("/admin")


@app.get("/")
def home(request: Request):
    require_device(request)
    return redirect("/pack")


# ---------- packing station ----------

@app.get("/pack", response_class=HTMLResponse)
def pack(request: Request, q: str = "", focus: str = ""):
    require_device(request)
    q = q.strip()
    with db.pool.connection() as conn:
        if q.isdigit() and len(q) == 13:
            product = service.product_by_ean(conn, q)
            if product:
                return redirect(f"/pack/product/{product['id']}")
        products = service.find_products(conn, q) or service.suggest_products(conn, q, limit=50)
    return render(request, "pack.html", q=q, products=products, focus=focus)


@app.get("/pack/suggest")
def pack_suggest(request: Request, q: str = ""):
    require_device(request)
    with db.pool.connection() as conn:
        products = service.suggest_products(conn, q.strip())
    return [{"id": p["id"], "article": p["article"], "size": p["size"], "color": p["color"],
             "stock": p["stock"]} for p in products]


@app.get("/pack/product/{product_id}", response_class=HTMLResponse)
def pack_product(request: Request, product_id: int):
    require_device(request)
    with db.pool.connection() as conn:
        product = service.get_product(conn, product_id)
        last_jobs = service.list_jobs(conn, query=product["ean"], limit=5) if product else []
    if not product:
        return redirect("/pack", request, "Товар не найден", error=True)
    return render(request, "product.html", product=product, last_jobs=last_jobs)


# Sticker PDFs live in memory for a few minutes under a random token: the packer sticker carries a personal code.
_stickers = {}


@app.post("/pack/sticker")
def pack_sticker(request: Request, code: str = Form("")):
    require_device(request)
    back = "/pack?focus=sticker"
    scanned = "".join(code.split())
    if not scanned:
        return redirect(back, request, "Отсканируйте баркод товара или личный код", error=True)
    key = auth.client_key(request)
    with db.pool.connection() as conn:
        row = nomenclature.by_barcode(conn, scanned) if scanned.isdigit() else None
        employee = None
        if not row:
            if auth.code_throttle.blocked(key):
                return redirect(back, request, "Слишком много неверных кодов. Подождите 15 минут.", error=True)
            employee = auth.user_by_code(conn, scanned)
        if row:
            pdf, title = stickers.product_pdf(conn, row), f"{row['article']} {row['size']}".strip()
        elif employee:
            auth.code_throttle.reset(key)
            pdf = stickers.packer_pdf(conn, employee["name"], auth.normalize_code(scanned),
                                      datetime.now(config.TIMEZONE).date())
            title = f"Проверила упаковала: {stickers.short_name(employee['name'])}"
        else:
            auth.code_throttle.fail(key)
            return redirect(back, request, f"Баркод {scanned} не найден в номенклатуре", error=True)
    now = time.time()
    for token in [t for t, v in _stickers.items() if now - v[1] > config.PRINT_LINK_MINUTES * 60]:
        _stickers.pop(token, None)
    token = secrets.token_urlsafe(16)
    _stickers[token] = (pdf, now, title)
    return redirect(f"/pack/sticker/{token}")


@app.get("/pack/sticker/{token}", response_class=HTMLResponse)
def pack_sticker_page(request: Request, token: str):
    require_device(request)
    item = _stickers.get(token)
    if not item:
        return redirect("/pack?focus=sticker", request, "Наклейка устарела, отсканируйте ещё раз", error=True)
    return render(request, "sticker.html", token=token, title=item[2])


@app.get("/pack/sticker/{token}/pdf")
def pack_sticker_pdf(request: Request, token: str):
    require_device(request)
    item = _stickers.get(token)
    if not item:
        return Response("Наклейка устарела, отсканируйте ещё раз.", status_code=404,
                        media_type="text/plain; charset=utf-8")
    return Response(item[0], media_type="application/pdf", headers={"Cache-Control": "no-store"})


def _employee(request, personal_code):
    if not personal_code or not personal_code.strip():
        return None, "Отсканируйте свой бейдж в поле «Личный код»."
    key = auth.client_key(request)
    if auth.code_throttle.blocked(key):
        return None, "Слишком много неверных личных кодов. Подождите 15 минут или обратитесь к администратору."
    with db.pool.connection() as conn:
        employee = auth.user_by_code(conn, personal_code)
    if not employee:
        auth.code_throttle.fail(key)
        return None, "Личный код не найден. Отсканируйте бейдж ещё раз."
    auth.code_throttle.reset(key)
    return employee, None


def _remember_job(request, job_id):
    jobs = [j for j in request.session.get("jobs", []) if time.time() - j[1] < config.PRINT_LINK_MINUTES * 60]
    jobs.append([job_id, time.time()])
    request.session["jobs"] = jobs[-20:]


def _job_allowed(request, job_id):
    return any(
        j[0] == job_id and time.time() - j[1] < config.PRINT_LINK_MINUTES * 60
        for j in request.session.get("jobs", [])
    )


@app.post("/print")
def do_print(request: Request, product_id: int = Form(...), qty: str = Form(""), personal_code: str = Form(""),
             size: str = Form(DEFAULT_SIZE)):
    require_device(request)
    back = f"/pack/product/{product_id}"
    count = _int(qty)
    if not count or count < 1:
        return redirect(back, request, "Укажите количество цифрами", error=True)
    employee, error = _employee(request, personal_code)
    if not employee:
        return redirect(back, request, error, error=True)
    try:
        job_id = service.print_codes(product_id, count, employee)
    except service.ServiceError as e:
        return redirect(back, request, str(e), error=True)
    _remember_job(request, job_id)
    return redirect(f"/jobs/{job_id}/print?size={_size(size)}")


@app.get("/jobs/{job_id}/print", response_class=HTMLResponse)
def job_print_page(request: Request, job_id: int, size: str = DEFAULT_SIZE):
    require_device(request)
    with db.pool.connection() as conn:
        job = service.get_job(conn, job_id)
    if not job:
        return redirect("/pack", request, "Печать не найдена", error=True)
    return render(request, "print.html", job=job, allowed=_job_allowed(request, job_id),
                  minutes=config.PRINT_LINK_MINUTES, size=_size(size))


@app.get("/jobs/{job_id}/pdf")
def job_pdf(request: Request, job_id: int, size: str = DEFAULT_SIZE):
    require_device(request)
    if not _job_allowed(request, job_id):
        return Response("Ссылка на печать устарела. Перепечатать может администратор из архива.",
                        status_code=403, media_type="text/plain; charset=utf-8")
    size = _size(size)
    with db.pool.connection() as conn:
        items = service.job_pdf_items(conn, job_id, size)
    buf = io.BytesIO()
    if size == "58x40":
        try:
            build_small_pdf(items, buf)
        except SmallLabelError as e:
            return Response(f"Этикетку 58×40 для этой печати собрать не удалось: {e}. Напечатайте размер 50×104.",
                            status_code=409, media_type="text/plain; charset=utf-8")
    else:
        build_print_pdf(items, buf)
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Cache-Control": "no-store",
                             "Content-Disposition": f'inline; filename="print-{job_id}-{size}.pdf"'})


# ---------- archive ----------

@app.get("/archive", response_class=HTMLResponse)
def archive(request: Request, q: str = "", code: str = "", user_id: str = "",
            date_from: str = "", date_to: str = ""):
    require_admin(request)
    with db.pool.connection() as conn:
        found_codes = service.find_code(conn, code) if code.strip() else None
        jobs = [] if found_codes is not None else service.list_jobs(
            conn, q, _int(user_id), _date(date_from), _date(date_to), limit=100
        )
        users = conn.execute("SELECT id, name FROM users WHERE code_hash IS NOT NULL ORDER BY name").fetchall()
    return render(request, "archive.html", q=q, code=code, user_id=_int(user_id),
                  date_from=date_from, date_to=date_to, jobs=jobs, found_codes=found_codes, users=users)


@app.get("/archive/job/{job_id}", response_class=HTMLResponse)
def archive_job(request: Request, job_id: int, highlight: int = 0):
    require_admin(request)
    with db.pool.connection() as conn:
        job = service.get_job(conn, job_id)
        codes = service.job_codes(conn, job_id) if job else []
    if not job:
        return redirect("/archive", request, "Печать не найдена", error=True)
    return render(request, "archive_job.html", job=job, codes=codes, highlight=highlight)


@app.post("/reprint")
async def reprint(request: Request):
    admin = require_admin(request)
    form = await request.form()
    job_id = _int(form.get("job_id"))
    back = f"/archive/job/{job_id}" if job_id else "/archive"
    code_ids = [int(v) for v in form.getlist("code_id") if str(v).isdigit()]
    try:
        new_job = service.reprint_codes(code_ids, admin)
    except service.ServiceError as e:
        return redirect(back, request, str(e), error=True)
    _remember_job(request, new_job)
    return redirect(f"/jobs/{new_job}/print?size={_size(form.get('size'))}")


# ---------- uploads ----------

@app.get("/uploads", response_class=HTMLResponse)
def uploads(request: Request):
    require_admin(request)
    with db.pool.connection() as conn:
        items = service.list_uploads(conn, limit=50)
    return render(request, "uploads.html", uploads=items)


@app.post("/uploads")
async def upload_files(request: Request, files: list[UploadFile] = File(...)):
    user = require_admin(request)
    ids = []
    for f in files:
        if not f.filename:
            continue
        if not f.filename.lower().endswith(".xlsx"):
            return redirect("/uploads", request, f"{f.filename}: нужен файл Excel (.xlsx)", error=True)
        content = await f.read()
        ids.append(service.create_upload(f.filename, content, user["id"]))
    if not ids:
        return redirect("/uploads", request, "Выберите файл", error=True)
    if len(ids) == 1:
        return redirect(f"/uploads/{ids[0]}")
    return redirect("/uploads", request, f"Файлов принято: {len(ids)}. Они обрабатываются по очереди.")


@app.get("/uploads/{upload_id}", response_class=HTMLResponse)
def upload_detail(request: Request, upload_id: int):
    require_admin(request)
    with db.pool.connection() as conn:
        upload = conn.execute(
            "SELECT up.*, u.name AS user_name FROM uploads up LEFT JOIN users u ON u.id = up.user_id WHERE up.id = %s",
            (upload_id,),
        ).fetchone()
        usage = service.upload_usage(conn, upload_id) if upload else None
        if upload and upload["deleted_by"]:
            upload["deleted_by_name"] = conn.execute(
                "SELECT name FROM users WHERE id = %s", (upload["deleted_by"],)).fetchone()["name"]
    if not upload:
        return redirect("/uploads", request, "Загрузка не найдена", error=True)
    return render(request, "upload.html", upload=upload, usage=usage)


@app.post("/uploads/{upload_id}/confirm")
def upload_confirm(request: Request, upload_id: int):
    require_admin(request)
    try:
        added = service.confirm_upload(upload_id)
    except service.ServiceError as e:
        return redirect(f"/uploads/{upload_id}", request, str(e), error=True)
    return redirect(f"/uploads/{upload_id}", request, f"Добавлено в хранилище: {added} шт.")


@app.post("/uploads/{upload_id}/cancel")
def upload_cancel(request: Request, upload_id: int):
    require_admin(request)
    service.cancel_upload(upload_id)
    return redirect("/uploads", request, "Загрузка отменена")


@app.post("/uploads/{upload_id}/delete")
def upload_delete(request: Request, upload_id: int, password: str = Form("")):
    admin = require_admin(request)
    with db.pool.connection() as conn:
        printed = service.upload_usage(conn, upload_id)["printed"]
    if printed:
        # Printed codes take their history with them, so the administrator confirms with the password.
        checked, error = _check_admin_password(request, password)
        if error or checked["id"] != admin["id"]:
            return redirect(f"/uploads/{upload_id}", request, (error or "Неверный пароль") + ". Загрузка не удалена.",
                            error=True)
    try:
        usage = service.delete_upload(upload_id, admin["id"], allow_printed=bool(printed))
    except service.ServiceError as e:
        return redirect(f"/uploads/{upload_id}", request, str(e), error=True)
    text = f"Загрузка №{upload_id} удалена: кодов из остатка {usage['stock']}"
    if usage["printed"]:
        text += f", напечатанных {usage['printed']}"
    return redirect("/uploads", request, text + ".")


# ---------- nomenclature ----------

def _nomenclature_page(request, q="", category="", edit=None, form=None, error=None):
    with db.pool.connection() as conn:
        rows, categories, loaded = nomenclature.listing(conn, q, category)
        editing = nomenclature.get_row(conn, edit) if edit else None
    if editing and not editing["manual"]:
        editing = None
    extra = {"flash": {"text": error, "error": True}} if error else {}
    return render(request, "nomenclature.html", rows=rows, categories=categories, loaded=loaded,
                  q=q, category=category, editing=editing, form=form or editing or {}, **extra)


@app.get("/nomenclature", response_class=HTMLResponse)
def nomenclature_page(request: Request, q: str = "", category: str = "", edit: int = 0):
    require_admin(request)
    return _nomenclature_page(request, q, category, edit or None)


@app.post("/nomenclature/add", response_class=HTMLResponse)
async def nomenclature_add(request: Request):
    admin = require_admin(request)
    form = dict(await request.form())
    try:
        item = nomenclature.save_manual(form, admin["name"])
    except nomenclature.NomenclatureError as e:
        return _nomenclature_page(request, form=form, error=str(e))
    return redirect("/nomenclature", request, f"Добавлено вручную: {item['article']} {item['size']}".strip())


@app.post("/nomenclature/{row_id}/edit", response_class=HTMLResponse)
async def nomenclature_edit(request: Request, row_id: int):
    admin = require_admin(request)
    form = dict(await request.form())
    try:
        item = nomenclature.save_manual(form, admin["name"], row_id)
    except nomenclature.NomenclatureError as e:
        return _nomenclature_page(request, edit=row_id, form=form, error=str(e))
    return redirect("/nomenclature", request, f"Строка изменена: {item['article']} {item['size']}".strip())


@app.post("/nomenclature/upload")
async def nomenclature_upload(request: Request, file: UploadFile = File(...)):
    admin = require_admin(request)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return redirect("/nomenclature", request, "Нужен файл Excel (.xlsx)", error=True)
    try:
        items = nomenclature.read_file(await file.read())
    except nomenclature.NomenclatureError as e:
        return redirect("/nomenclature", request, str(e), error=True)
    nomenclature.replace(items, file.filename, admin["name"])
    return redirect("/nomenclature", request, f"Номенклатура загружена: {len(items)} строк. Прежний список из файла заменён, строки «вручную» сохранены.")


# ---------- sticker designer ----------

SAMPLE_PRODUCT = {"category": "Рубашки", "article": "138N/белый", "size": "56", "barcode": "2039356108373"}
SAMPLE_PACKER = ("Михайлова Светлана", "0000000000")
PREVIEW_DPI = 304.8  # 12 px per mm


def _sample_data(conn, kind, sample):
    if kind == "packer":
        return stickers.packer_data(sample.strip() or SAMPLE_PACKER[0], SAMPLE_PACKER[1],
                                    datetime.now(config.TIMEZONE).date())
    row = nomenclature.by_barcode(conn, sample.strip()) if sample.strip().isdigit() else None
    if row is None and not sample.strip():
        row = conn.execute("SELECT * FROM nomenclature ORDER BY manual, position LIMIT 1").fetchone()
    return stickers.product_data(row or SAMPLE_PRODUCT)


def _kind(kind):
    return kind if kind in stickers.KINDS else "product"


@app.get("/designer", response_class=HTMLResponse)
def designer_page(request: Request, kind: str = "product"):
    require_admin(request)
    kind = _kind(kind)
    with db.pool.connection() as conn:
        layout = stickers.load(conn, kind)
    return render(request, "designer.html", kind=kind, kinds=stickers.KINDS, layout=layout,
                  fields=stickers.FIELDS[kind], code_type=stickers.CODES[kind], fonts=stickers.FONT_NAMES,
                  width=stickers.WIDTH_MM, height=stickers.HEIGHT_MM)


@app.post("/designer/preview")
async def designer_preview(request: Request):
    require_admin(request)
    try:
        body = await request.json()
        kind = _kind(body.get("kind"))
        layout = stickers.clean(kind, body.get("layout"))
    except (ValueError, AttributeError, stickers.LayoutError):
        return Response("Неверный макет", status_code=400, media_type="text/plain; charset=utf-8")
    with db.pool.connection() as conn:
        pdf = stickers.render(kind, layout, _sample_data(conn, kind, str(body.get("sample") or "")),
                              stickers.images_for(conn, layout))
    if body.get("format") == "pdf":
        return Response(pdf, media_type="application/pdf")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "s.pdf"
        src.write_bytes(pdf)
        subprocess.run(["pdftoppm", "-png", "-r", str(PREVIEW_DPI), "-singlefile", str(src), str(Path(tmp) / "s")],
                       check=True, timeout=30)
        png = (Path(tmp) / "s.png").read_bytes()
    return Response(png, media_type="image/png", headers={"Cache-Control": "no-store"})


@app.post("/designer/save")
async def designer_save(request: Request):
    require_admin(request)
    try:
        body = await request.json()
        kind = _kind(body.get("kind"))
        with db.tx() as conn:
            stickers.save(conn, kind, body.get("layout"))
    except (ValueError, AttributeError, stickers.LayoutError):
        return JSONResponse({"error": "Неверный макет"}, status_code=400)
    return {"ok": True}


@app.post("/designer/reset")
def designer_reset(request: Request, kind: str = Form("product")):
    require_admin(request)
    kind = _kind(kind)
    with db.tx() as conn:
        stickers.reset(conn, kind)
    return redirect(f"/designer?kind={kind}", request, f"Макет «{stickers.KINDS[kind]}» возвращён к исходному")


@app.post("/designer/image")
async def designer_image(request: Request, file: UploadFile = File(...)):
    require_admin(request)
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        return JSONResponse({"error": "Картинка больше 5 МБ"}, status_code=400)
    try:
        with db.tx() as conn:
            row = stickers.store_image(conn, content)
    except stickers.LayoutError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    return {"id": row["id"], "w": row["w"], "h": row["h"]}


# ---------- stock and journal ----------

@app.get("/stock", response_class=HTMLResponse)
def stock(request: Request, q: str = "", low: int = 0):
    require_admin(request)
    with db.pool.connection() as conn:
        rows = service.stock_report(conn)
        threshold = service.min_stock(conn)
    tokens = q.lower().split()
    rows = [r for r in rows if all(t in r["search"] for t in tokens)]
    if low:
        rows = [r for r in rows if r["stock"] < threshold]
    low_count = sum(1 for r in rows if r["stock"] < threshold)
    return render(request, "stock.html", rows=rows, q=q, low=low, threshold=threshold, low_count=low_count)


@app.post("/stock/threshold")
def stock_threshold(request: Request, value: str = Form(...)):
    require_admin(request)
    number = _int(value)
    if number is None or number < 0:
        return redirect("/stock", request, "Порог должен быть числом", error=True)
    with db.tx() as conn:
        db.set_setting(conn, "min_stock", number)
    return redirect("/stock", request, "Порог сохранён")


def _xlsx(headers, rows, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, len(h) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.get("/stock.xlsx")
def stock_xlsx(request: Request):
    require_admin(request)
    with db.pool.connection() as conn:
        rows = service.stock_report(conn)
    return _xlsx(
        ["Артикул", "Размер", "Цвет", "EAN", "В остатке", "В архиве"],
        [[r["article"], r["size"], r["color"], r["ean"], r["stock"], r["printed"]] for r in rows],
        "ostatki.xlsx",
    )


@app.get("/journal", response_class=HTMLResponse)
def journal(request: Request, q: str = "", user_id: str = "", date_from: str = "", date_to: str = "",
            kind: str = ""):
    require_admin(request)
    with db.pool.connection() as conn:
        jobs = service.list_jobs(conn, q, _int(user_id), _date(date_from), _date(date_to),
                                 kind or None, limit=500)
        uploads = [] if (q or user_id or kind) else service.list_uploads(conn, _date(date_from), _date(date_to))
        users = conn.execute("SELECT id, name FROM users WHERE code_hash IS NOT NULL ORDER BY name").fetchall()
    entries = [{"at": j["created_at"], "type": j["kind"], "row": j} for j in jobs]
    entries += [{"at": u["created_at"], "type": "upload", "row": u} for u in uploads]
    entries += [{"at": u["deleted_at"], "type": "upload_deleted", "row": u} for u in uploads if u["deleted_at"]]
    entries.sort(key=lambda e: e["at"], reverse=True)
    total_printed = sum(j["qty"] for j in jobs if j["kind"] == "print")
    total_reprinted = sum(j["qty"] for j in jobs if j["kind"] == "reprint")
    return render(request, "journal.html", entries=entries, q=q, user_id=_int(user_id),
                  date_from=date_from, date_to=date_to, kind=kind, users=users,
                  total_printed=total_printed, total_reprinted=total_reprinted)


@app.get("/journal.xlsx")
def journal_xlsx(request: Request, q: str = "", user_id: str = "", date_from: str = "", date_to: str = "",
                 kind: str = ""):
    require_admin(request)
    with db.pool.connection() as conn:
        jobs = service.list_jobs(conn, q, _int(user_id), _date(date_from), _date(date_to),
                                 kind or None, limit=100000)
    names = {"print": "Печать", "reprint": "Повторная печать"}
    return _xlsx(
        ["Дата и время", "Действие", "Сотрудник", "Артикул", "Размер", "EAN", "Количество", "Номер операции"],
        [[_local(j["created_at"]), names[j["kind"]], j["user_name"], j["article"] or "", j["size"] or "",
          j["ean"] or "", j["qty"], j["id"]] for j in jobs],
        "zhurnal.xlsx",
    )


# ---------- employees ----------

@app.get("/users", response_class=HTMLResponse)
def users_page(request: Request):
    require_admin(request)
    with db.pool.connection() as conn:
        users = conn.execute(
            "SELECT * FROM users WHERE role IN ('admin', 'packer')"
        ).fetchall()
    # «Упаковщик 2» before «Упаковщик 10»: numbers inside names sort by value.
    users.sort(key=lambda u: (not u["active"], u["role"] != "admin",
                              [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", u["name"])]))
    people = [{"id": u["id"], "name": u["name"], "code": u["personal_code"] or ""} for u in users]
    return render(request, "users.html", users=users, people=people, new_code=request.session.get("new_code"))


@app.get("/users/{user_id:int}", response_class=HTMLResponse)
def user_card(request: Request, user_id: int):
    require_admin(request)
    with db.pool.connection() as conn:
        person = conn.execute(
            "SELECT * FROM users WHERE id = %s AND role IN ('admin', 'packer')", (user_id,)
        ).fetchone()
    if not person:
        return redirect("/users", request, "Сотрудник не найден", error=True)
    new_code = request.session.get("new_code")
    return render(request, "user.html", person=person,
                  new_code=new_code if new_code and new_code["user_id"] == user_id else None)


def _valid_code(code):
    return len(auth.normalize_code(code)) >= 4


def _unique_error(e):
    constraint = getattr(e.diag, "constraint_name", "") or ""
    return "Такой личный код уже есть у другого сотрудника" if "code" in constraint else "Такой сотрудник уже есть"


def _password_taken(conn, password, except_id=None):
    admin = auth.admin_by_password(conn, password)
    return admin is not None and admin["id"] != except_id


@app.post("/users")
def create_user(request: Request, name: str = Form(...), role: str = Form(...),
                password: str = Form(""), personal_code: str = Form("")):
    require_admin(request)
    name, personal_code = name.strip(), personal_code.strip()
    if role not in auth.ROLE_NAMES or not name:
        return redirect("/users", request, "Укажите имя и роль", error=True)
    is_admin = role == "admin"
    if is_admin and len(password) < 8:
        return redirect("/users", request, "Для администратора нужен пароль не короче 8 символов", error=True)
    if personal_code and not _valid_code(personal_code):
        return redirect("/users", request, "Личный код должен быть не короче 4 символов", error=True)
    code = personal_code or auth.new_personal_code()
    try:
        with db.tx() as conn:
            if is_admin and _password_taken(conn, password):
                return redirect("/users", request, "Такой пароль уже у другого администратора, придумайте другой",
                                error=True)
            user_id = conn.execute(
                "INSERT INTO users (name, role, password_hash, code_hash, personal_code) "
                "VALUES (%s, %s, %s, %s, %s) RETURNING id",
                (name, role, auth.hash_password(password) if is_admin else None, auth.hash_code(code),
                 auth.normalize_code(code)),
            ).fetchone()["id"]
    except UniqueViolation as e:
        return redirect("/users", request, _unique_error(e), error=True)
    if not personal_code:
        request.session["new_code"] = {"user_id": user_id, "name": name, "code": code}
        return redirect(f"/users/{user_id}", request, f"Сотрудник «{name}» добавлен")
    return redirect("/users", request, f"Сотрудник «{name}» добавлен")


@app.post("/users/{user_id}/code")
def reset_code(request: Request, user_id: int, personal_code: str = Form("")):
    require_admin(request)
    card = f"/users/{user_id}"
    personal_code = personal_code.strip()
    if personal_code and not _valid_code(personal_code):
        return redirect(card, request, "Личный код должен быть не короче 4 символов", error=True)
    code = personal_code or auth.new_personal_code()
    try:
        with db.tx() as conn:
            row = conn.execute(
                "UPDATE users SET code_hash = %s, personal_code = %s WHERE id = %s RETURNING name",
                (auth.hash_code(code), auth.normalize_code(code), user_id),
            ).fetchone()
    except UniqueViolation as e:
        return redirect(card, request, _unique_error(e), error=True)
    if not row:
        return redirect("/users", request, "Сотрудник не найден", error=True)
    if personal_code:
        return redirect(card, request, "Код изменён, старый больше не действует")
    request.session["new_code"] = {"user_id": user_id, "name": row["name"], "code": code}
    return redirect(card, request, "Выдан новый личный код, старый больше не действует")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_people(filename, content):
    """Rows of (name, code) from an Excel or CSV file; a header row with «Имя»/«Код» is optional."""
    if filename.lower().endswith(".csv"):
        import csv
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            # Excel on Windows saves CSV in the Windows Cyrillic encoding.
            text = content.decode("cp1251", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        ws = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [[_cell_text(v) for v in r] for r in rows if any(_cell_text(v) for v in r)]
    name_col, code_col = 0, 1
    if rows:
        header = [h.lower() for h in rows[0]]
        names = [i for i, h in enumerate(header) if "имя" in h or "фио" in h or "сотрудник" in h]
        codes = [i for i, h in enumerate(header) if "код" in h]
        if names or codes:
            name_col = names[0] if names else 0
            code_col = codes[0] if codes else 1
            rows = rows[1:]
    return [(r[name_col] if name_col < len(r) else "", r[code_col] if code_col < len(r) else "") for r in rows]


@app.post("/users/import")
async def import_users(request: Request, file: UploadFile = File(...)):
    require_admin(request)
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".csv")):
        return redirect("/users", request, "Нужен файл Excel (.xlsx) или CSV", error=True)
    try:
        people = _read_people(file.filename, await file.read())
    except Exception:
        return redirect("/users", request, "Не удалось прочитать файл", error=True)
    added, updated, problems = 0, 0, []
    with db.tx() as conn:
        for line, (name, code) in enumerate(people, start=1):
            if not name or not _valid_code(code):
                problems.append(f"строка {line}: нет имени или код короче 4 символов")
                continue
            code_hash, shown = auth.hash_code(code), auth.normalize_code(code)
            owner = conn.execute("SELECT id, name FROM users WHERE code_hash = %s", (code_hash,)).fetchone()
            same_name = conn.execute(
                "SELECT id FROM users WHERE lower(name) = lower(%s) AND role IN ('admin', 'packer') "
                "ORDER BY active DESC, id LIMIT 1", (name,)
            ).fetchone()
            if owner and (not same_name or owner["id"] != same_name["id"]):
                if owner["name"].lower() != name.lower():
                    problems.append(f"{name}: этот код уже есть у «{owner['name']}»")
                    continue
                same_name = owner
            if same_name:
                # Loading the same list again refreshes codes of people already in the list.
                conn.execute("UPDATE users SET code_hash = %s, personal_code = %s WHERE id = %s",
                             (code_hash, shown, same_name["id"]))
                updated += 1
                continue
            conn.execute(
                "INSERT INTO users (name, role, code_hash, personal_code) VALUES (%s, 'packer', %s, %s)",
                (name, code_hash, shown),
            )
            added += 1
    message = f"Добавлено упаковщиков: {added}."
    if updated:
        message += f" Обновлены коды у уже записанных: {updated}."
    if problems:
        message += " Пропущено: " + "; ".join(problems[:10]) + ("…" if len(problems) > 10 else "")
    return redirect("/users", request, message, error=bool(problems))


@app.post("/users/{user_id}/toggle")
def toggle_user(request: Request, user_id: int, field: str = Form(...)):
    current = require_admin(request)
    card = f"/users/{user_id}"
    if field != "active":
        return redirect(card)
    if user_id == current["id"]:
        return redirect(card, request, "Нельзя отключить самого себя", error=True)
    with db.tx() as conn:
        row = conn.execute("UPDATE users SET active = NOT active WHERE id = %s RETURNING active",
                           (user_id,)).fetchone()
    return redirect(card, request, "Сотрудник включён" if row and row["active"] else "Сотрудник отключён")


@app.post("/users/{user_id}/edit")
def edit_user(request: Request, user_id: int, name: str = Form(""), role: str = Form(...), password: str = Form("")):
    current = require_admin(request)
    card = f"/users/{user_id}"
    name = name.strip()
    if not name or role not in auth.ROLE_NAMES:
        return redirect(card, request, "Укажите имя и роль", error=True)
    try:
        with db.tx() as conn:
            person = conn.execute("SELECT * FROM users WHERE id = %s AND role IN ('admin', 'packer') FOR UPDATE",
                                  (user_id,)).fetchone()
            if not person:
                return redirect("/users", request, "Сотрудник не найден", error=True)
            password_hash = person["password_hash"]
            if role == "admin" and person["role"] != "admin":
                if len(password) < 8:
                    return redirect(card, request, "Чтобы сделать администратором, задайте пароль не короче 8 символов",
                                    error=True)
                if _password_taken(conn, password, except_id=user_id):
                    return redirect(card, request, "Такой пароль уже у другого администратора, придумайте другой",
                                    error=True)
                password_hash = auth.hash_password(password)
            elif role == "packer" and person["role"] == "admin":
                if user_id == current["id"]:
                    return redirect(card, request, "Нельзя снять права администратора с самого себя", error=True)
                others = conn.execute("SELECT count(*) AS n FROM users WHERE role = 'admin' AND active AND id <> %s",
                                      (user_id,)).fetchone()["n"]
                if not others:
                    return redirect(card, request, "Это единственный администратор, его роль менять нельзя", error=True)
                password_hash = None
            conn.execute("UPDATE users SET name = %s, role = %s, password_hash = %s WHERE id = %s",
                         (name, role, password_hash, user_id))
    except UniqueViolation as e:
        return redirect(card, request, _unique_error(e), error=True)
    return redirect(card, request, "Данные сотрудника сохранены")


@app.post("/users/{user_id}/password")
def set_password(request: Request, user_id: int, password: str = Form(...)):
    require_admin(request)
    card = f"/users/{user_id}"
    if len(password) < 8:
        return redirect(card, request, "Пароль не короче 8 символов", error=True)
    with db.tx() as conn:
        if _password_taken(conn, password, except_id=user_id):
            return redirect(card, request, "Такой пароль уже у другого администратора, придумайте другой",
                            error=True)
        conn.execute("UPDATE users SET password_hash = %s WHERE id = %s AND role = 'admin'",
                     (auth.hash_password(password), user_id))
    return redirect(card, request, "Пароль изменён")


@app.get("/users/badge.pdf")
def badge(request: Request):
    require_admin(request)
    new_code = request.session.get("new_code")
    if not new_code:
        return redirect("/users", request, "Бейдж можно напечатать сразу после выдачи кода", error=True)
    return Response(badge_pdf(new_code["name"], new_code["code"]), media_type="application/pdf",
                    headers={"Cache-Control": "no-store"})


@app.post("/users/badge/done")
def badge_done(request: Request):
    require_admin(request)
    new_code = request.session.pop("new_code", None)
    return redirect(f"/users/{new_code['user_id']}" if new_code else "/users")


# ---------- computers ----------

@app.get("/devices", response_class=HTMLResponse)
def devices_page(request: Request):
    require_admin(request)
    with db.pool.connection() as conn:
        devices = conn.execute(
            "SELECT d.*, u.name AS admin_name FROM devices d LEFT JOIN users u ON u.id = d.approved_by "
            "ORDER BY d.active DESC, d.last_seen_at DESC NULLS LAST"
        ).fetchall()
    return render(request, "devices.html", devices=devices, this_id=request.state.device["id"])


@app.post("/devices/{device_id}/toggle")
def toggle_device(request: Request, device_id: int):
    require_admin(request)
    if device_id == request.state.device["id"]:
        return redirect("/devices", request, "Нельзя отключить компьютер, за которым вы сейчас работаете", error=True)
    with db.tx() as conn:
        conn.execute("UPDATE devices SET active = NOT active WHERE id = %s", (device_id,))
    return redirect("/devices")


@app.post("/devices/{device_id}/name")
def rename_device(request: Request, device_id: int, name: str = Form(...)):
    require_admin(request)
    if name.strip():
        with db.tx() as conn:
            conn.execute("UPDATE devices SET name = %s WHERE id = %s", (name.strip()[:60], device_id))
    return redirect("/devices")


@app.get("/health")
def health():
    with db.pool.connection() as conn:
        conn.execute("SELECT 1")
    return {"ok": True, "time": datetime.now(timezone.utc).isoformat()}
