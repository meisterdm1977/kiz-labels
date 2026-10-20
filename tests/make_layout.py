"""Builds a small label layout shaped like the 1С export: 36 rows per label, page break after each."""
import io
import random
import string

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.worksheet.pagebreak import Break
from PIL import Image
from pylibdmtx.pylibdmtx import encode

GS = "\x1d"
ROWS = 36


def make_code(gtin, rnd):
    alphabet = string.ascii_letters + string.digits + "!&?+-"
    serial = "".join(rnd.choice(alphabet) for _ in range(13))
    tail = "".join(rnd.choice(string.ascii_letters + string.digits) for _ in range(44))
    return f"01{gtin}21{serial}{GS}91EE12{GS}92{tail}", serial


def make_layout(path, count, article="Шорты 02 коричневый", size="40", ean="2047362903387",
                gtin="04610304896481", seed=1, broken=()):
    rnd = random.Random(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    for col, width in zip("ABCDE", (1, 5, 10, 5, 9)):
        ws.column_dimensions[col].width = width
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0
    codes = []
    for i in range(count):
        top = i * ROWS
        code, serial = make_code(gtin, rnd)
        codes.append(code)
        ws.cell(top + 4, 2, ean)
        ws.cell(top + 17, 2, "Артикул:")
        ws.cell(top + 17, 4, article)
        ws.cell(top + 19, 2, "Размер:")
        ws.cell(top + 19, 4, size)
        ws.cell(top + 21, 2, "Цвет:")
        ws.cell(top + 21, 4, "КОРИЧНЕВЫЙ")
        ws.cell(top + 35, 2, f"(01){gtin}\n(21){serial}")
        if i not in broken:
            enc = encode(code.encode("latin-1"))
            img = Image.frombytes("RGB", (enc.width, enc.height), enc.pixels)
            buf = io.BytesIO()
            img.save(buf, "PNG")
            buf.seek(0)
            xl = XlImage(buf)
            xl.width = xl.height = 90
            ws.add_image(xl, f"B{top + 26}")
        ws.row_breaks.append(Break(id=top + ROWS))
    wb.save(path)
    return codes
