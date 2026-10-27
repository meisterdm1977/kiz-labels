import re
import threading
import time

from pypdf import PdfReader

from kiz import db
from make_layout import make_layout


def approve(client, password="admin-password", name="Упаковка"):
    r = client.post("/device", data={"name": name, "password": password}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/pack"


def login(client, password="admin-password"):
    if not client.cookies.get("kiz_device"):
        approve(client, password)
    r = client.post("/admin", data={"password": password}, follow_redirects=False)
    assert r.status_code == 303 and "/admin" not in r.headers["location"]


def wait_upload(client, upload_id):
    for _ in range(240):
        with db.pool.connection() as conn:
            status = conn.execute("SELECT status FROM uploads WHERE id = %s", (upload_id,)).fetchone()["status"]
        if status != "processing":
            return status
        time.sleep(0.5)
    raise AssertionError("upload still processing")


def upload(client, path, name="Шорты 02 коричневый - 40-6шт.xlsx"):
    with open(path, "rb") as f:
        r = client.post("/uploads", files={"files": (name, f.read())}, follow_redirects=False)
    upload_id = int(r.headers["location"].rsplit("/", 1)[1])
    return upload_id, wait_upload(client, upload_id)


def create_packer(client, name):
    page = client.post("/users", data={"name": name, "role": "packer"}).text
    client.post("/users/badge/done")
    return re.search(r'class="mono big">(\d{8})<', page).group(1)


def stock_of(ean):
    with db.pool.connection() as conn:
        return conn.execute(
            "SELECT count(*) AS n FROM codes c JOIN products p ON p.id = c.product_id "
            "WHERE p.ean = %s AND c.status = 'stock'", (ean,)
        ).fetchone()["n"]


def test_full_cycle(client, tmp_path):
    assert client.get("/pack", follow_redirects=False).headers["location"] == "/device"
    login(client)
    anna = create_packer(client, "Анна")

    first = tmp_path / "a.xlsx"
    codes = make_layout(first, 6, seed=1)
    upload_id, status = upload(client, first)
    assert status == "ready"
    page = client.get(f"/uploads/{upload_id}").text
    assert "Новых: <b>6</b>" in page
    client.post(f"/uploads/{upload_id}/confirm")
    assert stock_of("2047362903387") == 6

    # The same file again: every code is a duplicate, nothing is added.
    upload_id2, _ = upload(client, first)
    assert "Новых: <b>0</b>" in client.get(f"/uploads/{upload_id2}").text
    client.post(f"/uploads/{upload_id2}/cancel")

    # A new batch of the same product adds to the stock.
    second = tmp_path / "b.xlsx"
    make_layout(second, 3, seed=2)
    upload_id3, _ = upload(client, second)
    client.post(f"/uploads/{upload_id3}/confirm")
    assert stock_of("2047362903387") == 9

    # Scanning the EAN opens the product, text search finds it by article and size.
    r = client.get("/pack", params={"q": "2047362903387"}, follow_redirects=False)
    product_url = r.headers["location"]
    product_id = int(product_url.rsplit("/", 1)[1])
    assert "Шорты 02 коричневый" in client.get("/pack", params={"q": "шорты 02 40"}).text

    # Wrong personal code prints nothing.
    client.post("/print", data={"product_id": product_id, "qty": "2", "personal_code": "000"})
    assert stock_of("2047362903387") == 9

    r = client.post("/print", data={"product_id": product_id, "qty": "4", "personal_code": anna},
                    follow_redirects=False)
    job_id = int(re.search(r"/jobs/(\d+)/print", r.headers["location"]).group(1))
    assert stock_of("2047362903387") == 5

    pdf = client.get(f"/jobs/{job_id}/pdf")
    assert pdf.status_code == 200
    pages = PdfReader(__import__("io").BytesIO(pdf.content)).pages
    assert len(pages) == 4
    assert round(float(pages[0].mediabox.height) / 72 * 25.4) == 104
    small = PdfReader(__import__("io").BytesIO(client.get(f"/jobs/{job_id}/pdf", params={"size": "58x40"}).content)).pages
    assert len(small) == 4
    assert [round(float(v) / 72 * 25.4) for v in (small[0].mediabox.width, small[0].mediabox.height)] == [58, 40]
    assert "58×40" in client.get(f"/jobs/{job_id}/print", params={"size": "58x40"}).text

    # Oldest codes go first.
    with db.pool.connection() as conn:
        printed = [r["code"] for r in conn.execute(
            "SELECT c.code FROM print_job_codes jc JOIN codes c ON c.id = jc.code_id "
            "WHERE jc.job_id = %s ORDER BY jc.position", (job_id,))]
    assert printed == codes[:4]

    # Cannot print more than the stock.
    client.post("/print", data={"product_id": product_id, "qty": "6", "personal_code": anna})
    assert stock_of("2047362903387") == 5

    # Archive: find a code by its scan without GS separators, the administrator reprints it.
    scan = codes[1].replace("\x1d", "")
    archive = client.get("/archive", params={"code": scan}).text
    assert "напечатан" in archive
    with db.pool.connection() as conn:
        code_id = conn.execute("SELECT id FROM codes WHERE code = %s", (codes[1],)).fetchone()["id"]
    r = client.post("/reprint", data={"job_id": job_id, "code_id": code_id}, follow_redirects=False)
    reprint_job = int(re.search(r"/jobs/(\d+)/print", r.headers["location"]).group(1))
    assert len(PdfReader(__import__("io").BytesIO(client.get(f"/jobs/{reprint_job}/pdf").content)).pages) == 1
    assert stock_of("2047362903387") == 5

    journal = client.get("/journal").text
    assert "Анна" in journal and "Повторная печать" in journal and "Администратор" in journal
    assert "Напечатано: <b>4</b>" in journal and "Перепечатано: <b>1</b>" in journal
    assert client.get("/journal.xlsx").status_code == 200
    assert client.get("/stock.xlsx").status_code == 200


def test_print_link_expires_for_other_sessions(client):
    with db.pool.connection() as conn:
        job = conn.execute("SELECT id FROM print_jobs ORDER BY id LIMIT 1").fetchone()
    from fastapi.testclient import TestClient
    from kiz.main import app
    other = TestClient(app)
    login(other)
    assert other.get(f"/jobs/{job['id']}/pdf").status_code == 403


def test_parallel_printing_never_shares_codes(client, tmp_path):
    path = tmp_path / "c.xlsx"
    make_layout(path, 20, ean="2047000000001", size="42", seed=3)
    upload_id, _ = upload(client, path)
    client.post(f"/uploads/{upload_id}/confirm")
    from kiz import service
    with db.pool.connection() as conn:
        product_id = conn.execute("SELECT id FROM products WHERE ean = '2047000000001'").fetchone()["id"]
        user = conn.execute("SELECT * FROM users WHERE role = 'packer' LIMIT 1").fetchone()

    results, errors = [], []

    def worker():
        try:
            results.append(service.print_codes(product_id, 3, user))
        except service.ServiceError as e:
            errors.append(str(e))

    threads = [threading.Thread(target=worker) for _ in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    with db.pool.connection() as conn:
        rows = conn.execute(
            "SELECT jc.code_id FROM print_job_codes jc WHERE jc.job_id = ANY(%s)", (results,)
        ).fetchall()
    ids = [r["code_id"] for r in rows]
    assert len(ids) == len(set(ids)) == 3 * len(results)
    assert len(results) == 6 and len(errors) == 2
    assert stock_of("2047000000001") == 2


def test_password_guessing_is_blocked():
    from fastapi.testclient import TestClient
    from kiz import auth
    from kiz.main import app

    auth.login_throttle.failures.clear()
    c = TestClient(app)
    for _ in range(10):
        c.post("/device", data={"password": "wrong"})
    r = c.post("/device", data={"password": "admin-password"}, follow_redirects=False)
    assert r.headers["location"] == "/device"
    auth.login_throttle.failures.clear()
    login(c)


def test_own_and_imported_personal_codes(client, tmp_path):
    import openpyxl
    login(client)
    client.post("/users", data={"name": "Олег", "role": "packer", "personal_code": "0042-77"})
    with db.pool.connection() as conn:
        from kiz import auth
        assert auth.user_by_code(conn, "0042-77")["name"] == "Олег"

    # The same code twice is refused.
    page = client.post("/users", data={"name": "Иван", "role": "packer", "personal_code": "0042-77"}).text
    assert "уже есть у другого сотрудника" in page

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ФИО", "Код"])
    ws.append(["Мария", "00123"])
    ws.append(["Пётр", 98765])
    ws.append(["Дубль", "0042-77"])
    path = tmp_path / "people.xlsx"
    wb.save(path)
    page = client.post("/users/import", files={"file": ("people.xlsx", path.read_bytes())}).text
    assert "Добавлено упаковщиков: 2" in page and "Олег" in page
    with db.pool.connection() as conn:
        assert auth.user_by_code(conn, "00123")["name"] == "Мария"
        assert auth.user_by_code(conn, "98765")["name"] == "Пётр"

        petr_id = auth.user_by_code(conn, "98765")["id"]
    client.post(f"/users/{petr_id}/code", data={"personal_code": "5555"})
    with db.pool.connection() as conn:
        assert auth.user_by_code(conn, "98765") is None
        assert auth.user_by_code(conn, "5555")["name"] == "Пётр"


def test_only_approved_computers_and_admin_sections(client):
    from fastapi.testclient import TestClient
    from kiz import auth, config
    from kiz.main import app

    stranger = TestClient(app)
    for path in ["/pack", "/archive", "/stock", "/users", "/journal.xlsx", "/pack/product/1"]:
        assert stranger.get(path, follow_redirects=False).headers["location"] == "/device"
    # A made-up device cookie does not help.
    stranger.cookies.set("kiz_device", "guess")
    assert stranger.get("/pack", follow_redirects=False).headers["location"] == "/device"

    packing = TestClient(app)
    approve(packing, name="Упаковка 2")
    assert packing.get("/pack").status_code == 200
    for path in ["/archive", "/uploads", "/stock", "/journal", "/users", "/devices"]:
        r = packing.get(path, follow_redirects=False)
        assert r.headers["location"].startswith("/admin?next="), path
    assert "/admin" in packing.post("/reprint", data={"code_id": 1}, follow_redirects=False).headers["location"]
    assert "Архив" not in packing.get("/pack").text and "Администратор" in packing.get("/pack").text

    # Administrator mode closes after the idle time.
    login(packing)
    assert packing.get("/stock").status_code == 200
    config.ADMIN_IDLE_MINUTES, saved = 0, config.ADMIN_IDLE_MINUTES
    try:
        assert packing.get("/stock", follow_redirects=False).status_code == 303
    finally:
        config.ADMIN_IDLE_MINUTES = saved
    login(packing)
    packing.post("/admin/logout")
    assert packing.get("/stock", follow_redirects=False).status_code == 303

    # The administrator switches a computer off and it loses access at once.
    login(client)
    with db.pool.connection() as conn:
        device_id = conn.execute("SELECT id FROM devices WHERE name = 'Упаковка 2'").fetchone()["id"]
    assert "Упаковка 2" in client.get("/devices").text
    client.post(f"/devices/{device_id}/toggle")
    assert packing.get("/pack", follow_redirects=False).headers["location"] == "/device"

    # Administrators sign in by password alone, so two of them cannot share one.
    page = client.post("/users", data={"name": "Второй", "role": "admin", "password": "admin-password"}).text
    assert "уже у другого администратора" in page
    client.post("/users", data={"name": "Второй", "role": "admin", "password": "second-pass"})
    other = TestClient(app)
    login(other, "second-pass")
    auth.login_throttle.failures.clear()


def test_staff_list_shows_codes_and_reimport_updates(client, tmp_path):
    import openpyxl
    login(client)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Имя", "Код"])
    ws.append(["Зоя", "7701"])
    path = tmp_path / "zoya.xlsx"
    wb.save(path)
    client.post("/users/import", files={"file": ("zoya.xlsx", path.read_bytes())})
    with db.pool.connection() as conn:
        zoya = conn.execute("SELECT * FROM users WHERE name = 'Зоя'").fetchone()
        # A packer added before codes were kept readable: only the hash is known.
        conn.execute("UPDATE users SET personal_code = NULL WHERE id = %s", (zoya["id"],))
    assert "7701" not in client.get("/users").text

    # Loading the same list again fills in the code instead of reporting a duplicate.
    page = client.post("/users/import", files={"file": ("zoya.xlsx", path.read_bytes())}).text
    assert "Обновлены коды у уже записанных: 1" in page and ">7701<" in page
    with db.pool.connection() as conn:
        assert conn.execute("SELECT count(*) AS n FROM users WHERE name = 'Зоя'").fetchone()["n"] == 1

    # The list is read-only; changes happen on the employee card.
    card = client.get(f"/users/{zoya['id']}").text
    assert "Отключить сотрудника" in card and "Сменить код" in card
    page = client.post(f"/users/{zoya['id']}/toggle", data={"field": "active"}).text
    assert "Сотрудник отключён" in page and "Включить сотрудника" in page
    page = client.post(f"/users/{zoya['id']}/code", data={"personal_code": "7702"}).text
    assert "Код изменён" in page and "7702" in page
    assert client.get("/users/badge.pdf", follow_redirects=False).status_code in (200, 303)


def test_product_suggestions_forgive_typos(client):
    login(client)
    found = client.get("/pack/suggest", params={"q": "шотрты 40"}).json()
    assert found and found[0]["article"] == "Шорты 02 коричневый" and found[0]["size"] == "40"
    assert "stock" in found[0]
    assert client.get("/pack/suggest", params={"q": "платье"}).json() == []
    # The search page falls back to the same typo-tolerant match.
    assert "Шорты 02 коричневый" in client.get("/pack", params={"q": "шотрты"}).text
    from fastapi.testclient import TestClient
    from kiz.main import app
    assert TestClient(app).get("/pack/suggest", params={"q": "шорты"}, follow_redirects=False).status_code == 303


def test_deleting_uploads(client, tmp_path):
    login(client)
    anna = create_packer(client, "Борис")
    path = tmp_path / "a.xlsx"
    make_layout(path, 4, ean="2047362900011", seed=7)
    printed_upload, _ = upload(client, path)
    client.post(f"/uploads/{printed_upload}/confirm")
    product_id = int(client.get("/pack", params={"q": "2047362900011"}, follow_redirects=False)
                     .headers["location"].rsplit("/", 1)[1])
    client.post("/print", data={"product_id": product_id, "qty": "1", "personal_code": anna})

    other = tmp_path / "b.xlsx"
    make_layout(other, 3, ean="2047362900028", size="42", seed=8)
    fresh_upload, _ = upload(client, other, name="Шорты 02 коричневый - 42-3шт.xlsx")
    client.post(f"/uploads/{fresh_upload}/confirm")
    assert stock_of("2047362900028") == 3

    # Nothing printed: a plain confirmation is enough, the product leaves the search.
    assert "Удалить загрузку?" not in client.get(f"/uploads/{printed_upload}").text
    r = client.post(f"/uploads/{fresh_upload}/delete", follow_redirects=False)
    assert r.headers["location"] == "/uploads"
    with db.pool.connection() as conn:
        assert conn.execute("SELECT count(*) AS n FROM codes WHERE upload_id = %s", (fresh_upload,)).fetchone()["n"] == 0
        assert not conn.execute("SELECT 1 FROM products WHERE ean = '2047362900028'").fetchone()

    # Printed codes: the page warns, and only the administrator's password deletes them with their history.
    page = client.get(f"/uploads/{printed_upload}").text
    assert "Будет удалено 1 напечатанных кодов" in page
    client.post(f"/uploads/{printed_upload}/delete", data={"password": "wrong"})
    assert stock_of("2047362900011") == 3
    client.post(f"/uploads/{printed_upload}/delete", data={"password": "admin-password"})
    with db.pool.connection() as conn:
        assert conn.execute("SELECT count(*) AS n FROM codes WHERE upload_id = %s", (printed_upload,)).fetchone()["n"] == 0
        assert conn.execute("SELECT count(*) AS n FROM print_jobs WHERE product_id = %s", (product_id,)).fetchone()["n"] == 0
        assert conn.execute("SELECT status FROM uploads WHERE id = %s", (printed_upload,)).fetchone()["status"] == "deleted"
    journal = client.get("/journal").text
    assert journal.count("Удаление загрузки") == 2
    assert "удалено" in client.get("/uploads").text


def test_side_menu_shows_sections_by_role(client):
    approve(client)
    client.post("/admin/logout")
    page = client.get("/pack").text
    assert 'class="panel"' in page and "Администратор" in page and "/journal" not in page
    login(client)
    page = client.get("/journal").text
    assert all(f'href="{url}"' in page for url in ("/pack", "/archive", "/uploads", "/stock", "/journal", "/users", "/devices"))
    assert 'class="on"><svg' in page and "Выйти" in page


def test_nomenclature_upload_replaces_the_list(client, tmp_path):
    import openpyxl
    login(client)
    path = tmp_path / "nom.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Категория", "Артикул/цвет", "Размер", "Основной баркод"])
    ws.append(["Брюки", "2409-2 бежевый", "M", "2043747204766"])
    ws.append(["Шорты", "Шорты 02 коричневый", "2XL", 2047687243885])
    wb.save(path)
    with open(path, "rb") as f:
        client.post("/nomenclature/upload", files={"file": ("nom.xlsx", f.read())})
    page = client.get("/nomenclature").text
    assert "2409-2 бежевый" in page and "2047687243885" in page and "Строк: 2" in page
    assert "Строк: 1" in client.get("/nomenclature", params={"category": "Шорты"}).text
    assert "Строк: 1" in client.get("/nomenclature", params={"q": "бежев"}).text

    ws.delete_rows(3)
    wb.save(path)
    with open(path, "rb") as f:
        client.post("/nomenclature/upload", files={"file": ("nom2.xlsx", f.read())})
    page = client.get("/nomenclature").text
    assert "Шорты 02" not in page and "Строк: 1" in page and "nom2.xlsx" in page

    bad = client.post("/nomenclature/upload", files={"file": ("x.txt", b"hello")}, follow_redirects=True).text
    assert "Нужен файл Excel" in bad and "Строк: 1" in bad


def test_nomenclature_manual_rows_survive_uploads(client, tmp_path):
    import openpyxl
    login(client)

    def upload(rows, name):
        wb = openpyxl.Workbook()
        wb.active.append(["Категория", "Артикул/цвет", "Размер", "Основной баркод"])
        for r in rows:
            wb.active.append(r)
        path = tmp_path / name
        wb.save(path)
        client.post("/nomenclature/upload", files={"file": (name, path.read_bytes())})

    upload([["Брюки", "Файловые брюки", "M", "2000000000011"]], "f1.xlsx")
    added = client.post("/nomenclature/add", data={"category": "Юбки", "article": "Ручная юбка",
                                                   "size": "S", "barcode": "2000000000028"}, follow_redirects=True).text
    assert "Добавлено вручную: Ручная юбка S" in added and "вручную</span>" in added and "Строк: 2" in added

    dup = client.post("/nomenclature/add", data={"article": "Дубль", "barcode": "2000000000011"}).text
    assert "уже есть: Файловые брюки, размер M" in dup and 'value="Дубль"' in dup

    page = client.get("/nomenclature", params={"q": "Ручная"}).text
    row_id = int(re.search(r"edit=(\d+)", page).group(1))
    edited = client.post(f"/nomenclature/{row_id}/edit", data={"category": "Юбки", "article": "Ручная юбка 2",
                                                              "size": "M", "barcode": "2000000000028"},
                         follow_redirects=True).text
    assert "Строка изменена: Ручная юбка 2 M" in edited

    upload([["Брюки", "Другие брюки", "L", "2000000000035"]], "f2.xlsx")
    page = client.get("/nomenclature").text
    assert "Файловые брюки" not in page and "Ручная юбка 2" in page and "Строк: 2" in page

    upload([["Юбки", "Юбка из файла", "M", "2000000000028"]], "f3.xlsx")
    page = client.get("/nomenclature").text
    assert "Ручная юбка 2" not in page and "Юбка из файла" in page and "Строк: 1" in page


def test_barcode_and_packer_stickers(client):
    import io
    login(client)
    client.post("/nomenclature/add", data={"category": "Рубашки", "article": "777N/белый", "size": "56",
                                           "barcode": "2039356100001"})
    code = create_packer(client, "Смирнова Ольга")
    page = client.get("/pack").text
    assert "Печатать КИЗ" in page and "Печатать наклейку ШК" in page

    r = client.post("/pack/sticker", data={"code": "2039356100001"}, follow_redirects=True)
    assert "777N/белый 56" in r.text
    token = re.search(r"/pack/sticker/([\w-]+)/pdf", r.text).group(1)
    pdf = client.get(f"/pack/sticker/{token}/pdf").content
    page = PdfReader(io.BytesIO(pdf)).pages[0]
    assert round(float(page.mediabox.width) / 72 * 25.4) == 58 and round(float(page.mediabox.height) / 72 * 25.4) == 40
    text = page.extract_text()
    assert "Рубашка" in text and "Артикул: 777N" in text and "Цвет: БЕЛЫЙ" in text and "ИП Малоземов" in text
    assert "2039356100001" in text

    r = client.post("/pack/sticker", data={"code": code}, follow_redirects=True)
    assert "Проверила упаковала: Ольга С" in r.text
    token = re.search(r"/pack/sticker/([\w-]+)/pdf", r.text).group(1)
    pdf = client.get(f"/pack/sticker/{token}/pdf").content
    assert "Ольга С" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text()

    r = client.post("/pack/sticker", data={"code": "2000000000999"}, follow_redirects=True)
    assert "не найден в номенклатуре" in r.text


def test_sticker_designer_layout_is_used_for_printing(client):
    import io
    login(client)
    page = client.get("/designer").text
    assert "Конструктор наклеек 58×40" in page and "Проверила упаковала" in page

    from kiz import stickers
    layout = stickers.DEFAULTS["product"] + [
        {"id": "mine", "type": "text", "x": 2, "y": 2, "w": 40, "h": 5, "text": "Спасибо за покупку", "field": None,
         "font": "Arial", "size": 9, "bold": False, "align": "center"},
        {"id": "l", "type": "line", "x": 2, "y": 20, "w": 54, "h": 0.3},
        {"id": "bad", "type": "script"},
    ]
    png = client.post("/designer/preview", json={"kind": "product", "layout": layout}).content
    assert png[:4] == b"\x89PNG"
    assert client.post("/designer/save", json={"kind": "product", "layout": layout}).json() == {"ok": True}

    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "black").save(buf, "PNG")
    img = client.post("/designer/image", files={"file": ("logo.png", buf.getvalue())}).json()
    assert img["w"] == 40 and img["h"] == 20
    assert "error" in client.post("/designer/image", files={"file": ("x.png", b"nope")}).json()

    client.post("/nomenclature/add", data={"category": "Платья", "article": "888 синий", "size": "M",
                                           "barcode": "2039356100002"})
    r = client.post("/pack/sticker", data={"code": "2039356100002"}, follow_redirects=True)
    token = re.search(r"/pack/sticker/([\w-]+)/pdf", r.text).group(1)
    text = PdfReader(io.BytesIO(client.get(f"/pack/sticker/{token}/pdf").content)).pages[0].extract_text()
    assert "Спасибо за покупку" in text and "Платье" in text and "Цвет: СИНИЙ" in text

    client.post("/designer/reset", data={"kind": "product"})
    r = client.post("/pack/sticker", data={"code": "2039356100002"}, follow_redirects=True)
    token = re.search(r"/pack/sticker/([\w-]+)/pdf", r.text).group(1)
    assert "Спасибо" not in PdfReader(io.BytesIO(client.get(f"/pack/sticker/{token}/pdf").content)).pages[0].extract_text()
