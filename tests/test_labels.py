from pypdf import PdfReader

from kiz.labels import code_prefix, read_workbook, render_pdf
from make_layout import make_layout


def test_reads_codes_and_fields(tmp_path):
    path = tmp_path / "layout.xlsx"
    codes = make_layout(path, 3)
    labels = read_workbook(path)
    assert [l.code for l in labels] == codes
    assert all(not l.errors for l in labels)
    first = labels[0]
    assert (first.article, first.size, first.ean, first.color) == (
        "Шорты 02 коричневый", "40", "2047362903387", "КОРИЧНЕВЫЙ")
    assert first.code_prefix == code_prefix(codes[0]) and len(first.code_prefix) == 31


def test_unreadable_datamatrix_is_reported(tmp_path):
    path = tmp_path / "layout.xlsx"
    make_layout(path, 3, broken={1})
    labels = read_workbook(path)
    assert labels[1].errors == ["не удалось прочитать DataMatrix"]
    assert not labels[0].errors and not labels[2].errors


def test_renders_one_label_sized_page_per_code(tmp_path):
    path = tmp_path / "layout.xlsx"
    make_layout(path, 4)
    out = tmp_path / "labels.pdf"
    render_pdf(path, out, 4)
    pages = PdfReader(out).pages
    assert len(pages) == 4
    box = pages[0].mediabox
    assert round(float(box.width) / 72 * 25.4) == 50
    assert round(float(box.height) / 72 * 25.4) == 104


def test_small_label_keeps_the_datamatrix_and_is_built_for_old_uploads(tmp_path):
    from kiz import small_label
    path = tmp_path / "source.xlsx"
    codes = make_layout(path, 3)
    labels = small_label.load(tmp_path)  # no small.json yet: made from source.xlsx
    assert (tmp_path / small_label.FILE).exists()
    assert [small_label.modules_decode(l["modules"]) for l in labels] == codes
    assert labels[0]["ean"] == "2047362903387"
    assert ["Артикул:", "Шорты 02 коричневый"] in labels[0]["fields"]
    out = tmp_path / "small.pdf"
    small_label.build_pdf([(tmp_path, 2), (tmp_path, 0)], str(out))
    pages = PdfReader(out).pages
    assert len(pages) == 2
    assert [round(float(v) / 72 * 25.4) for v in (pages[0].mediabox.width, pages[0].mediabox.height)] == [58, 40]


def test_long_texts_shrink_instead_of_being_cut(tmp_path):
    import openpyxl
    from kiz.labels import _fit_texts
    path = tmp_path / "layout.xlsx"
    make_layout(path, 1, article="Шорты женские летние хлопковые с поясом 02 коричневый")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.merge_cells("D17:E17")
    _fit_texts(wb)
    assert 4.5 <= ws["D17"].font.sz < 11  # the long article is made smaller to fit
    assert ws["D21"].font.sz == 11  # a short text keeps its size


def test_long_field_values_take_the_whole_width(tmp_path):
    import openpyxl
    from kiz.labels import _widen_fields
    path = tmp_path / "layout.xlsx"
    make_layout(path, 1, article="Костюм спортивный женский оверсайз 1234 тёмно-зелёный")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for label, value in (("B17:C17", "D17:E17"), ("B19:C19", "D19:E19")):
        ws.merge_cells(label)
        ws.merge_cells(value)
    _widen_fields(wb)
    assert ws["B17"].value == "Артикул: Костюм спортивный женский оверсайз 1234 тёмно-зелёный"
    assert ws["B19"].value == "Размер: 40" and ws["D19"].value is None
    assert "B17:E17" in [str(r) for r in ws.merged_cells.ranges]
